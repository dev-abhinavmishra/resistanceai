"""
Builds ResistanceAI_Presentation_v9.pptx for the Congressional App Challenge video.

Written for a general audience: no oncology background assumed.
Slide titles are full sentences; detail lives in speaker notes.

Layout system: title zone ends at 2.05in, content runs 2.15-6.85in,
footer carries source lines and the page number. Content sits in cards
and columns rather than floating text.

Run:  python deck/build_deck.py
Deps: pip install python-pptx pillow lxml
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from PIL import Image
from pathlib import Path

# ---------- theme: 2 fonts, 1 dark + 1 light + 1 accent ----------
# Ink + teal: keeps the app's dark medical feel but trades the flat blue
# for a warmer, handwritten-lab-notebook character.
BG     = RGBColor(0x0A, 0x12, 0x19)
SURF   = RGBColor(0x13, 0x20, 0x2B)
LINE   = RGBColor(0x23, 0x3A, 0x46)
TXT    = RGBColor(0xF2, 0xF7, 0xF6)
MUT    = RGBColor(0x8C, 0xA3, 0xAD)
ACC    = RGBColor(0x14, 0xB8, 0xA6)
ACCTXT = RGBColor(0x2D, 0xD4, 0xBF)
RING   = RGBColor(0x15, 0x4A, 0x44)
BTNTXT = RGBColor(0x06, 0x2B, 0x27)
CODEBG = RGBColor(0x06, 0x0C, 0x11)

UI   = "Georgia"     # serif voice for titles and prose
MONO = "Consolas"    # monospace for code, numbers, lab tags

DEMO_URL = "https://resistanceai.streamlit.app"
TOTAL    = 14

ROOT = Path(__file__).resolve().parents[1]
SHOT = ROOT / "docs" / "screenshots"
OUT  = ROOT / "ResistanceAI_Presentation_v9.pptx"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def box(s, x, y, w, h, fill=SURF, line=None, round_=False):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            shp.adjustments[0] = 0.05
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(s, x, y, w, h, runs, size=24, bold=False, color=TXT, font=UI,
         align=PP_ALIGN.LEFT, spacing=1.08, sp_after=8, anchor=MSO_ANCHOR.TOP,
         wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [runs]
    first = True
    for run in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(sp_after)
        if isinstance(run, tuple):
            t, kw = run
        else:
            t, kw = run, {}
        r = p.add_run()
        r.text = t
        f = r.font
        f.size = Pt(kw.get("size", size))
        f.bold = kw.get("bold", bold)
        f.color.rgb = kw.get("color", color)
        f.name = kw.get("font", font)
    return tb


def title(s, txt, sub=None):
    box(s, 0.78, 0.42, 0.55, 0.05, fill=ACC)
    text(s, 0.75, 0.62, 11.85, 1.35, txt, size=40, bold=True, spacing=1.0)
    if sub:
        text(s, 0.78, 1.62, 11.6, 0.5, sub, size=20, color=MUT)


def eyebrow(s, x, y, txt, size=13):
    box(s, x, y + 0.07, 0.12, 0.12, fill=ACC)
    text(s, x + 0.28, y, min(7.0, 12.95 - x), 0.4, txt, size=size,
         bold=True, color=ACCTXT, font=MONO)


def linkify(tb, url=DEMO_URL):
    for p in tb.text_frame.paragraphs:
        for r in p.runs:
            r.hyperlink.address = url
            r.font.underline = True


def pageno(s, n):
    text(s, 12.4, 7.1, 0.85, 0.3, f"{n:02d} / {TOTAL}", size=11, color=MUT,
         font=MONO, align=PP_ALIGN.RIGHT, sp_after=0)


def ring(s, x, y, d, width=1.5):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                             Inches(d), Inches(d))
    shp.fill.background()
    shp.line.color.rgb = RING
    shp.line.width = Pt(width)
    shp.shadow.inherit = False


def src_line(s, txt, y=7.05):
    text(s, 0.75, y, 11.85, 0.35, txt, size=12, color=MUT)


def pic(s, path, x, y, w, line=True):
    with Image.open(path) as im:
        ih = w * im.height / im.width
    s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    if line:
        box(s, x, y, w, ih, fill=None, line=LINE)
    return ih


def stat(s, x, y, w, big, small, h=1.25):
    box(s, x, y, w, h, fill=SURF, line=LINE, round_=True)
    text(s, x + 0.25, y + 0.14, w - 0.5, 0.55, big, size=30, bold=True,
         font=MONO, color=ACCTXT)
    text(s, x + 0.25, y + 0.78, w - 0.4, 0.4, small, size=13, color=MUT,
         font=MONO)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# ================= 1. TITLE =================
s = slide()
box(s, 0, 0, 13.333, 0.09, fill=ACC)
ring(s, 10.7, -1.9, 5.0)
ring(s, 11.9, -0.7, 2.6)
eyebrow(s, 0.78, 1.7, "CONGRESSIONAL APP CHALLENGE", size=15)
text(s, 0.75, 2.35, 11.8, 1.4, "ResistanceAI", size=66, bold=True)
text(s, 0.78, 3.75, 10.8, 1.2,
     "Software that helps cancer doctors choose treatments "
     "less likely to fail.", size=30, color=MUT, spacing=1.15)
text(s, 0.78, 5.9, 11.8, 0.5,
     "Abhinav Mishra     Python and Streamlit     trained on free public data",
     size=18, color=MUT)
linkify(text(s, 0.78, 6.55, 11.8, 0.5,
             "Open the live demo  ›  resistanceai.streamlit.app",
             size=18, color=ACCTXT, font=MONO))
notes(s, "Hi, I'm Abhinav. Doctors who treat cancer face a hard problem: even "
         "with a tumor's DNA report in hand, it is hard to know which drug will "
         "actually work. I built a small app that estimates that up front. "
         "In about three minutes I will show you the whole thing, live.")

# ================= 2. WHAT CANCER IS =================
s = slide()
pageno(s, 2)
title(s, "Cancer starts when DNA instructions go wrong.")
text(s, 0.85, 2.3, 6.9, 4.3, [
    "A tumor's DNA is its damaged instruction manual.",
    "Sequencing machines can now read that damage.",
    "One tumor can carry thousands of errors; a few drive it.",
    "Some drivers have drugs built to block them.",
    "Matching the error to the drug is precision oncology.",
], size=24, sp_after=16)
box(s, 8.25, 2.15, 4.35, 4.5, round_=True)
eyebrow(s, 8.6, 2.55, "WORDS TO KNOW")
gloss = [("MUTATION", "a DNA error, found by sequencing"),
         ("DRIVER", "an error that pushes the cancer"),
         ("TARGETED DRUG", "built to block one specific error"),
         ("PRECISION ONCOLOGY", "matching the drug to the error")]
y = 3.15
for term, dfn in gloss:
    text(s, 8.6, y, 3.7, 0.35, term, size=14, bold=True, color=ACCTXT,
         font=MONO)
    text(s, 8.6, y + 0.38, 3.7, 0.45, dfn, size=16, color=MUT)
    y += 0.88
notes(s, "A quick primer for anyone without a biology background. Cancer begins "
         "when a cell's DNA is damaged and the cell grows out of control. "
         "Clinics now sequence a tumor's DNA, and for many of the errors they "
         "find, called mutations, there is a drug designed to hit that exact "
         "weakness. The glossary on the right gives you every word you need "
         "for the rest of this deck.")

# ================= 3. THE GAP =================
s = slide()
pageno(s, 3)
title(s, "Knowing which drug will work takes experts most hospitals lack.")
text(s, 0.85, 2.25, 6.9, 3.2, [
    "One report can list hundreds of mutations.",
    "Hospital panels read hundreds of genes at once.",
    "Specialist tumor boards live at big centers.",
    "Community oncologists often decide alone.",
    "A wrong first choice can cost months.",
], size=24, sp_after=14)
box(s, 0.85, 5.55, 0.07, 0.85, fill=ACC)
text(s, 1.15, 5.68, 6.6, 0.9,
     "Sequencing is routine. Reading it is the bottleneck.",
     size=20, color=ACCTXT, spacing=1.15)
box(s, 8.25, 2.15, 4.35, 2.1, round_=True)
eyebrow(s, 8.6, 2.5, "WHAT AN ONCOLOGIST IS")
text(s, 8.6, 3.05, 3.7, 1.0,
     "The doctor who treats cancer and chooses the drug.",
     size=17, color=MUT, spacing=1.15)
box(s, 8.25, 4.45, 4.35, 2.2, round_=True)
eyebrow(s, 8.6, 4.8, "WHAT A TUMOR BOARD IS")
text(s, 8.6, 5.35, 3.7, 1.1,
     "A panel of specialists who review hard reports together.",
     size=17, color=MUT, spacing=1.15)
notes(s, "When a tumor report comes back, a panel of specialists called a "
         "molecular tumor board can review it and recommend a drug. Those "
         "panels mostly exist at large academic cancer centers. The oncologist, "
         "the cancer doctor treating most patients, is often reading that "
         "report alone under time pressure, and a drug that cannot work still "
         "costs the patient weeks on the clock.")

# ================= 4. BIG NUMBER =================
s = slide()
pageno(s, 4)
title(s, "So the model learned from real patients, not rules I wrote.")
text(s, 0.75, 2.2, 7.4, 1.6, "26,064", size=105, bold=True, font=MONO,
     color=ACCTXT)
text(s, 0.8, 3.95, 7.0, 1.1,
     "cancer patients with a sequenced tumor and an outcome on record.",
     size=24, color=TXT, spacing=1.15)
box(s, 8.3, 2.15, 4.25, 3.25, round_=True)
eyebrow(s, 8.65, 2.5, "WHERE IT CAME FROM")
text(s, 8.65, 3.05, 3.6, 0.9, [
    "10,336   MSK-IMPACT, 2017",
    "24,950   MSK-CHORD, 2024",
], size=15, font=MONO, color=TXT, sp_after=8)
text(s, 8.65, 4.05, 3.6, 1.1,
     "Two public datasets from Memorial Sloan Kettering.",
     size=16, color=MUT, spacing=1.2)
for i, (big, small) in enumerate([
        ("15,638", "PATIENTS TO TRAIN"),
        ("5,213", "TO CALIBRATE"),
        ("5,213", "HELD BACK FOR TESTING")]):
    stat(s, 0.8 + i * 4.15, 5.55, 3.9, big, small)
src_line(s, "Source: MSK-IMPACT (Zehir et al., Nat Med 2017) and MSK-CHORD "
            "(Nature 2024) via cBioPortal.")
notes(s, "Instead of writing rules by hand, I trained the model on two public "
         "datasets from Memorial Sloan Kettering Cancer Center: MSK-IMPACT "
         "and MSK-CHORD. After cleaning, 26,064 patients had both a sequenced "
         "tumor and a recorded outcome. They were split three ways: fifteen "
         "thousand to train on, five thousand to calibrate the score, and "
         "five thousand the model never touched, reserved for the report card "
         "you will see later.")

# ================= 5. THE PRODUCT =================
s = slide()
pageno(s, 5)
title(s, "It estimates the chance the next treatment fails.")
pic(s, str(SHOT / "v9_risk_full.png"), 0.7, 2.2, 7.55)
text(s, 0.8, 6.05, 7.4, 0.4,
     "A real lung cancer case, loaded from the worklist.",
     size=15, color=MUT)
box(s, 8.55, 2.2, 4.0, 3.3, round_=True)
eyebrow(s, 8.9, 2.55, "WHAT YOU ARE SEEING")
text(s, 8.9, 3.1, 3.4, 2.3, [
    "a 53 percent risk score",
    "an honest uncertainty band",
    "32nd percentile in the corpus",
    "2 mechanisms flagged by name",
], size=17, sp_after=12)
eyebrow(s, 8.55, 5.85, "TRY IT LIVE")
linkify(text(s, 8.55, 6.3, 4.0, 0.5, "resistanceai.streamlit.app",
             size=16, color=ACCTXT, font=MONO))
notes(s, "Here it is running. I load a real lung cancer patient and the app "
         "scores a 53 percent chance the next treatment fails, places that "
         "patient at the 32nd percentile of everyone on record, and names two "
         "biological escape routes it found. Notice the band beside the score: "
         "the model is honest that this profile is unusual, and unusual is "
         "exactly when a doctor should slow down. The link at the bottom is "
         "the public demo.")

# ================= 6. IT SHOWS ITS WORK =================
s = slide()
pageno(s, 6)
title(s, "It shows exactly which evidence moved the number.")
pic(s, str(SHOT / "v9_attrib.png"), 0.7, 2.15, 7.9)
text(s, 0.8, 6.4, 7.6, 0.4,
     "Left: what moved the score. Right: the cited evidence.",
     size=15, color=MUT)
box(s, 8.8, 2.15, 3.75, 4.35, round_=True)
eyebrow(s, 9.1, 2.5, "WHAT THE CARDS CITE")
text(s, 9.1, 3.05, 3.2, 3.2, [
    "STK11/KEAP1 cold tumor warning",
    "HRD context: PARP drug sensitivity",
    "MSI or POLE: checkpoint response",
    "each names the mutation it found",
], size=16, sp_after=14)
notes(s, "Judges and doctors ask the same question: why this number. The chart "
         "on the left removes each feature one at a time and measures how much "
         "the score moves. The cards on the right name biological escape "
         "routes the tumor might use, and each card cites the specific "
         "mutation it found in this patient's DNA.")

# ================= 7. WHAT TO TRY NEXT =================
s = slide()
pageno(s, 7)
title(s, "Then it lines up what is worth trying next.")
pic(s, str(SHOT / "v9_therapies.png"), 0.75, 2.15, 5.7)
pic(s, str(SHOT / "v9_twins.png"),     6.9, 2.15, 5.7)
eyebrow(s, 0.8, 5.55, "DRUGS + TRIALS")
text(s, 0.8, 6.0, 5.5, 0.9,
     "Regimens matched to this tumor's mutations; recruiting trials "
     "arrive live from ClinicalTrials.gov.",
     size=16, color=MUT, spacing=1.15)
eyebrow(s, 6.95, 5.55, "SIMILAR PATIENTS")
text(s, 6.95, 6.0, 5.5, 0.9,
     "What happened to the 8 nearest real patients: survival and "
     "failure rate.",
     size=16, color=MUT, spacing=1.15)
notes(s, "A risk score alone does not help anyone act. So the app maps the "
         "patient's mutations to treatment regimens with their evidence "
         "grades, searches ClinicalTrials.gov live for recruiting studies, "
         "and shows what actually happened to the eight most similar patients "
         "in the data.")

# ================= 8. HOW IT WORKS (diagram) =================
s = slide()
pageno(s, 8)
title(s, "Behind the screen, four steps a doctor can inspect.")
steps = [("DNA report", "sequencing file"),
         ("201 signals", "mutations and history"),
         ("8 models vote", "each sees a slice"),
         ("Probability", "calibrated, with range")]
bx, bw, bh = 0.85, 2.55, 1.95
for i, (t1, t2) in enumerate(steps):
    x = bx + i * 3.13
    box(s, x, 2.45, bw, bh, round_=True)
    text(s, x, 2.62, bw, 0.6, str(i + 1), size=28, bold=True,
         color=ACCTXT, font=MONO, align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 3.3, bw - 0.3, 0.55, t1, size=22, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, x + 0.15, 3.88, bw - 0.3, 0.45, t2, size=14, color=MUT,
         align=PP_ALIGN.CENTER)
    if i < 3:
        text(s, x + bw + 0.06, 3.05, 0.6, 0.7, "›", size=36,
             color=ACCTXT, align=PP_ALIGN.CENTER)
box(s, 0.85, 5.15, 5.7, 1.65, round_=True)
eyebrow(s, 1.15, 5.4, "THE 201 SIGNALS")
text(s, 1.15, 5.85, 5.1, 0.8,
     "Mutations, gene copy changes, TMB, MSI status, stage, smoking, purity.",
     size=16, color=MUT, spacing=1.15)
box(s, 6.8, 5.15, 5.7, 1.65, round_=True)
eyebrow(s, 7.1, 5.4, "WHY EIGHT MODELS")
text(s, 7.1, 5.85, 5.1, 0.8,
     "Each trains on a different slice; their spread measures doubt.",
     size=16, color=MUT, spacing=1.15)
notes(s, "Technically the pipeline is four steps. The DNA report is cleaned "
         "into a fixed list of 201 signals covering mutations, copy number "
         "changes, tumor mutation burden, MSI status, stage, smoking and "
         "purity. Eight gradient boosting models, each trained on a different "
         "slice of patients, each make a prediction. A final calibration step "
         "turns their average vote into a true probability, and how much they "
         "disagree becomes the honest band on screen.")

# ================= 9. DATA-FLOW MODEL =================
s = slide()
pageno(s, 9)
title(s, "Trace the data: trained once, then scored per request.")

# DFD shapes: circles = processes, cylinders = data stores,
# rectangles = external entities, arrows = data flows (dashed = load/read).


def node_lbl(xc, y, t, w=1.1, size=9):
    text(s, xc - w / 2, y, w, 0.3, t, size=size, color=MUT, font=MONO,
         align=PP_ALIGN.CENTER, sp_after=0)


def flow(x1, y1, x2, y2, dashed=False, color=None):
    conn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                  Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color or (MUT if dashed else ACC)
    conn.line.width = Pt(1.5 if dashed else 1.75)
    ln = conn.line._get_or_add_ln()
    if dashed:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    ln.append(ln.makeelement(
        qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    conn.shadow.inherit = False


def entity(x, y, w, name, sub=None, h=0.85):
    box(s, x, y, w, h)
    text(s, x + 0.08, y + 0.14, w - 0.16, 0.35, name, size=11.5, bold=True,
         align=PP_ALIGN.CENTER, sp_after=0)
    if sub:
        text(s, x + 0.08, y + 0.5, w - 0.16, 0.28, sub, size=9, color=MUT,
             font=MONO, align=PP_ALIGN.CENTER, sp_after=0)


def process(xc, yc, name, sub, d=1.2):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(xc - d / 2),
                             Inches(yc - d / 2), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = ACC
    shp.line.color.rgb = ACC
    shp.shadow.inherit = False
    text(s, xc - d / 2, yc - 0.3, d, 0.32, name, size=12, bold=True,
         color=BTNTXT, align=PP_ALIGN.CENTER, sp_after=0)
    text(s, xc - d / 2, yc + 0.05, d, 0.26, sub, size=8, color=BTNTXT,
         font=MONO, align=PP_ALIGN.CENTER, sp_after=0)


def store(x, y, w, name, sub, h=0.9):
    shp = s.shapes.add_shape(MSO_SHAPE.CAN, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = SURF
    shp.line.color.rgb = ACC
    shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    text(s, x, y + h * 0.34, w, 0.3, name, size=12, bold=True,
         align=PP_ALIGN.CENTER, sp_after=0)
    text(s, x, y + h * 0.62, w, 0.26, sub, size=8, color=MUT, font=MONO,
         align=PP_ALIGN.CENTER, sp_after=0)


# ---- lane 1: offline training ----
eyebrow(s, 0.9, 2.12, "OFFLINE — TRAINS ONCE")
entity(0.85, 2.72, 1.5, "cBioPortal", "IMPACT + CHORD")
process(3.6, 3.15, "Ingest", "ingest_v2.py")
store(4.85, 2.7, 1.4, "Corpus", "26,064 records")
process(7.5, 3.15, "Train", "8 HGB + iso")
store(8.75, 2.7, 1.4, "Artifact", "pkl + metrics")
box(s, 10.45, 2.7, 2.05, 0.9)
text(s, 10.6, 2.82, 1.75, 0.7,
     "re-run to retrain — weights stay local, mounted at deploy",
     size=11, color=MUT, spacing=1.05, sp_after=0)
flow(2.35, 3.15, 2.98, 3.15)
flow(4.22, 3.15, 4.82, 3.15)
flow(6.28, 3.15, 6.88, 3.15)
flow(8.12, 3.15, 8.72, 3.15)
node_lbl(2.67, 2.82, "raw data")
node_lbl(4.52, 2.82, "records")
node_lbl(6.58, 2.82, "features")
node_lbl(8.42, 2.82, "weights")

# ---- lane 2: live request ----
eyebrow(s, 0.9, 4.32, "LIVE — EVERY REQUEST")
entity(0.85, 4.98, 1.5, "Genomic report", "upload · worklist")
process(3.6, 5.4, "Normalize", "features.py")
process(5.42, 5.4, "Predict", "predictor.py")
process(7.24, 5.4, "Match", "twins · cohorts")
entity(8.75, 4.98, 1.5, "Streamlit UI", "8 tabs · Flask")
entity(10.75, 4.98, 1.5, "Oncologist", "reads, decides")
flow(2.35, 5.4, 2.98, 5.4)
flow(4.22, 5.4, 4.8, 5.4)
flow(6.04, 5.4, 6.62, 5.4)
flow(7.86, 5.4, 8.72, 5.4)
flow(10.28, 5.4, 10.72, 5.4)
node_lbl(2.67, 5.07, "report")
node_lbl(4.51, 5.07, "signals")
node_lbl(6.33, 5.07, "proba")
node_lbl(8.29, 5.07, "matches")
node_lbl(10.5, 5.07, "read", w=0.5)

# store-to-process loads at app start (dashed, one teal one gray so
# their crossing reads cleanly)
flow(9.45, 3.62, 5.68, 4.76, dashed=True, color=ACCTXT)
flow(5.55, 3.62, 7.0, 4.76, dashed=True)
node_lbl(8.35, 4.28, "weights", w=0.9)
node_lbl(5.75, 4.28, "corpus", w=0.9)

# live external service
box(s, 8.5, 6.3, 2.1, 0.55)
text(s, 8.5, 6.34, 2.1, 0.24, "ClinicalTrials.gov", size=11, bold=True,
     align=PP_ALIGN.CENTER, sp_after=0)
text(s, 8.5, 6.6, 2.1, 0.22, "live trial search", size=8, color=MUT,
     font=MONO, align=PP_ALIGN.CENTER, sp_after=0)
flow(8.85, 6.28, 7.62, 6.02)

# legend
LY = 6.55
shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), Inches(LY),
                         Inches(0.2), Inches(0.2))
shp.fill.solid(); shp.fill.fore_color.rgb = ACC
shp.line.color.rgb = ACC; shp.shadow.inherit = False
text(s, 1.15, LY - 0.04, 0.85, 0.3, "process", size=9, color=MUT,
     font=MONO, sp_after=0)
box(s, 2.1, LY, 0.2, 0.2)
text(s, 2.4, LY - 0.04, 0.85, 0.3, "entity", size=9, color=MUT, font=MONO,
     sp_after=0)
shp = s.shapes.add_shape(MSO_SHAPE.CAN, Inches(3.35), Inches(LY - 0.03),
                         Inches(0.22), Inches(0.28))
shp.fill.solid(); shp.fill.fore_color.rgb = SURF
shp.line.color.rgb = ACC; shp.shadow.inherit = False
text(s, 3.68, LY - 0.04, 0.8, 0.3, "store", size=9, color=MUT, font=MONO,
     sp_after=0)
flow(4.6, LY + 0.1, 4.98, LY + 0.1)
text(s, 5.08, LY - 0.04, 1.0, 0.3, "data flow", size=9, color=MUT,
     font=MONO, sp_after=0)
flow(6.15, LY + 0.1, 6.53, LY + 0.1, dashed=True)
text(s, 6.63, LY - 0.04, 1.4, 0.3, "load / fetch", size=9, color=MUT,
     font=MONO, sp_after=0)
notes(s, "Read it like a data-flow diagram. Circles are things the code "
         "does, cylinders are stored data, rectangles are the outside "
         "world, and every arrow is labeled with what moves along it. "
         "Top lane, run once: data leaves cBioPortal, ingest cleans it "
         "into the 201-signal schema, the corpus stores 26,064 records, "
         "training writes one calibrated artifact. Bottom lane, every "
         "request: a genomic report is normalized, Predict turns it into "
         "a score, Match finds similar patients and safe cohorts, and the "
         "result reaches the oncologist through Streamlit or the Flask "
         "API. The two dashed lines are the only crossings: the artifact "
         "loads into Predict and the corpus loads into Match at startup. "
         "ClinicalTrials.gov is the one live fetch — trials are searched "
         "fresh for each patient.")

# ================= 10. CODE =================
s = slide()
pageno(s, 10)
title(s, "It even scores how sure it is.")
code = [
    ("# eight models each make a call",                            MUT),
    ("votes = [m.predict_proba(x)[:, 1][0] for m in self.models]", TXT),
    ("proba = self.isotonic.predict([np.mean(votes)])[0]",         TXT),
    ("score = int(round(proba * 100))    # the estimate",          TXT),
    ("",                                                           TXT),
    ("# disagreement becomes the band on screen",                  MUT),
    ("half  = max(2, int(round(1.645 * np.std(votes) * 100)))",    TXT),
    ("ci    = [max(0, score - half), min(100, score + half)]",     TXT),
]
box(s, 0.85, 2.15, 6.7, 4.15, fill=CODEBG, line=LINE, round_=True)
y = 2.35
for t, c in code:
    if t:
        text(s, 1.1, y, 6.2, 0.42, t, size=13, font=MONO, color=c,
             sp_after=0)
    y += 0.45
text(s, 0.9, 6.5, 6.5, 0.5, "backend/predictor.py", size=13, color=MUT,
     font=MONO)
text(s, 8.05, 2.3, 4.5, 4.4, [
    "Eight models vote; their average becomes the score.",
    "Their disagreement sets the band's width.",
    "A wide band means slow down and look closer.",
    "Calibrated on 5,213 patients it never saw.",
], size=24, sp_after=22)
notes(s, "This is the honest part, in the actual code. The eight models vote, "
         "their average is calibrated into a probability, and how much they "
         "disagree sets the width of the band on screen. It is a disagreement "
         "measure, not a guarantee, and the app never pretends to be sure "
         "when it is not.")

# ================= 11. READABLE CODE =================
s = slide()
pageno(s, 11)
title(s, "Some of the most careful code is the easiest to read.")
panels = [
    (0.85, "FINDING SIMILAR PATIENTS", "backend/digital_twin.py",
     "Closest biology wins: real neighbors, with their real outcomes.",
     [("# how different each tumor is from this one",        MUT),
      ("diff = (self._mat[idx] - qv) / self._std",           TXT),
      ("dist = np.sqrt((diff ** 2).mean(axis=1))",           TXT),
      ("order = np.argsort(dist)[:k]   # the k closest",     TXT)]),
    (6.84, "PROTECTING SMALL COHORTS", "backend/cohort_analytics.py",
     "Under five matches, it reports counts and keeps details private.",
     [("MIN_COHORT_FOR_STATS = 5",                           TXT),
      ("if n < MIN_COHORT_FOR_STATS:",                       TXT),
      ("    # report the count, hide the details",           MUT),
      ("    return {**base, \"suppressed_small_n\": True}",  TXT)]),
]
for px, label, fname, cap, lines in panels:
    eyebrow(s, px + 0.05, 2.15, label)
    box(s, px, 2.5, 5.64, 2.6, fill=CODEBG, line=LINE, round_=True)
    y = 2.75
    for t, c in lines:
        if t:
            text(s, px + 0.28, y, 5.2, 0.35, t, size=14, font=MONO,
                 color=c, sp_after=0)
        y += 0.42
    text(s, px + 0.05, 5.3, 5.4, 0.4, fname, size=13, color=MUT, font=MONO)
    text(s, px + 0.05, 5.85, 5.5, 1.0, cap, size=24, color=TXT)
notes(s, "Two more spots worth a read. The twin engine measures how "
         "different every patient's biology is and keeps the closest "
         "neighbors, so the similar patients you saw earlier are real "
         "matches. And in the analytics code, when a question matches fewer "
         "than five patients it returns only the count, so no small cohort "
         "can be picked apart.")

# ================= 12. THE REPORT CARD =================
s = slide()
pageno(s, 12)
title(s, "Tested on patients it never saw, the numbers hold.")
pic(s, str(SHOT / "chart_calibration.png"), 0.8, 2.15, 5.9)
text(s, 7.15, 2.3, 5.35, 2.7, [
    ("When it said 60%, about 60% failed.", {"size": 24}),
    ("Graded on 5,213 patients it never trained on.", {"size": 24}),
    ("AUC 0.78 overall; 0.73 on the 2017 set.", {"size": 24}),
], sp_after=20)
box(s, 7.15, 5.35, 5.35, 1.5, round_=True)
eyebrow(s, 7.45, 5.55, "AUC BY CANCER TYPE")
text(s, 7.45, 5.95, 5.0, 0.9, [
    "breast      0.81    lung        0.77",
    "colorectal  0.74    prostate    0.74",
    "pancreatic  0.71",
], size=13, color=TXT, font=MONO, sp_after=3)
src_line(s, "Source: backend/artifacts/model_metrics.json. "
            "ROC AUC 0.78, calibration error 0.017.")
notes(s, "Here is the report card on 5,213 patients the model never trained "
         "on. Each dot asks the same question: when the model predicted a "
         "number, what share of patients actually failed. The dots sit near "
         "the diagonal, which is what honest probability looks like. "
         "Performance holds across cancer types, from 0.81 on breast cancer "
         "down to 0.71 on pancreatic, and across both source datasets. It is "
         "not a crystal ball, and the app says that on the screen.")

# ================= 13. CLOSE =================
s = slide()
pageno(s, 13)
box(s, 0, 7.35, 13.333, 0.15, fill=ACC)
ring(s, 10.9, 4.4, 4.4)
ring(s, 12.0, 5.5, 2.2)
text(s, 0.75, 1.8, 11.8, 1.3,
     "Doctors should not have to guess alone.", size=44, bold=True)
text(s, 0.78, 3.3, 11.5, 1.6, [
    "Free public data. Open code. Real patients.",
    "Watch it live: load a demo patient and judge for yourself.",
], size=26, color=MUT, sp_after=14)
text(s, 0.78, 4.9, 11.5, 0.5,
     "26,064 patients  ·  201 signals  ·  8 models  ·  0.78 AUC",
     size=16, color=MUT, font=MONO)
btn = box(s, 0.78, 5.6, 4.5, 0.85, fill=ACC, round_=True)
btn.click_action.hyperlink.address = DEMO_URL
btb = text(s, 0.78, 5.6, 4.5, 0.85, [
    ("CLICK TO OPEN", {"size": 10, "color": BTNTXT, "font": MONO,
                       "bold": True}),
    ("resistanceai.streamlit.app", {"size": 19, "color": BTNTXT,
                                    "font": MONO, "bold": True}),
], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=2)
btb.click_action.hyperlink.address = DEMO_URL
notes(s, "That is ResistanceAI. It gives a doctor the kind of read an expert "
         "panel would give, in seconds, with the reasoning left on the "
         "screen. Everything you saw is the running app on real public data, "
         "and the button on screen opens the public demo at "
         "resistanceai.streamlit.app. Thanks for watching.")

# ================= 14. SOURCES =================
s = slide()
pageno(s, 14)
title(s, "Sources")
text(s, 0.85, 2.1, 11.6, 4.7, [
    "Zehir et al. MSK-IMPACT cohort, Nature Medicine 2017, via cBioPortal.",
    "MSK-CHORD cohort, Nature 2024, PMID 39506116, via cBioPortal.",
    "ClinicalTrials.gov API v2, US National Library of Medicine.",
    "Evaluation: backend/artifacts/model_metrics.json, test n = 5,213.",
    "Model code: backend/predictor.py, digital_twin.py, cohort_analytics.py.",
    "Screenshots captured from the running app.",
    "This deck was generated by deck/build_deck.py.",
], size=22, color=MUT, sp_after=14)
notes(s, "Every number in this deck traces to one of these sources or to the "
         "app's own evaluation artifacts in the repository.")

# PowerPoint renders hyperlink runs in the theme's hlink color, overriding
# our run color — retint the theme so links stay on-palette.
from lxml import etree
from pptx.oxml.ns import qn

for part in prs.part.package.iter_parts():
    if "theme" in str(part.partname):
        root = etree.fromstring(part.blob)
        scheme = root.find(".//" + qn("a:clrScheme"))
        for tag in ("a:hlink", "a:folHlink"):
            el = scheme.find(qn(tag))
            if el is None:
                continue
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, qn("a:srgbClr")).set("val", "2DD4BF")
        part._blob = etree.tostring(root, xml_declaration=True,
                                    encoding="UTF-8", standalone=True)

prs.save(OUT)
print(f"saved {OUT}  |  {len(prs.slides._sldIdLst)} slides")
